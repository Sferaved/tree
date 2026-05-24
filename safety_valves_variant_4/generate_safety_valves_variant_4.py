from pathlib import Path
import math

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUT_DIR = Path(__file__).resolve().parent
DOCX_PATH = OUT_DIR / "Praktychna_robota_zapobizhni_klapany_variant_4.docx"
XLSX_PATH = OUT_DIR / "rozrahunky_zapobizhni_klapany_variant_4.xlsx"
MD_PATH = OUT_DIR / "Praktychna_robota_zapobizhni_klapany_variant_4.md"


def fmt(value, digits=2):
    return f"{value:.{digits}f}".replace(".", ",")


def setup_doc(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(14)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)


def paragraph(doc, text="", bold=False, align="justify", first_line=True):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    if first_line and align == "justify":
        p.paragraph_format.first_line_indent = Cm(1.25)
    run = p.add_run(str(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.bold = bold
    return p


def heading(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text.upper())
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.bold = True
    return p


def set_cell(cell, text, bold=False, center=True):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade(cell):
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "D9EAF7")
    props.append(shd)


def table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement(f"w:{name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")
        borders.append(border)
    tbl_pr.append(borders)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(table)
    for i, header in enumerate(headers):
        set_cell(table.rows[0].cells[i], header, bold=True)
        shade(table.rows[0].cells[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value)
    return table


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.italic = True


# Variant 4 source data and calculations.
G_BOILER = 1300  # kg/h
P_BOILER_MPA = 1.64
P_BOILER_PA = P_BOILER_MPA * 1_000_000
H_LIFT = 0.40  # cm
D_PLATE = 4.0  # cm
H_RATIO = H_LIFT / D_PLATE
K_VALVE = 0.015
N_CALC = G_BOILER / (K_VALVE * P_BOILER_PA * D_PLATE * H_LIFT)
N_ACCEPTED = 2

P1R = 26  # kgf/cm2
P2 = 2.4  # kgf/cm2
G_AIR = 21  # kg/h
RHO = 14.8  # kg/m3
ADIABATIC_K = 1.4
ALPHA = 0.6
P1 = 1.15 * P1R
P_RATIO = P2 / P1
B_COEFF = 0.505
F_NEEDED = G_AIR / (0.00159 * ALPHA * B_COEFF * math.sqrt(RHO * (P1 - P2)))
SEAT_DIAMETER = 60
F_SELECTED = 2240
INLET_DIAMETER = 100


CONTROL_ANSWERS = [
    ("Дайте визначення посудин, що працюють під тиском.",
     "Посудини, що працюють під тиском, - це герметично закриті ємності, призначені для здійснення технологічних процесів, зберігання або транспортування газоподібних, рідких та інших речовин під надлишковим тиском."),
    ("На які посудини розповсюджуються правила охорони праці під час експлуатації обладнання, що працює під тиском?",
     "Правила поширюються на посудини з водою температурою вище 110 °C, посудини з парою або газом тиском понад 0,5 бар, балони, цистерни, бочки для зріджених і стиснутих газів, посудини для спорожнення під тиском та барокамери."),
    ("З яким тиском балони для транспортування і зберігання газів відносяться до обладнання з підвищеною небезпекою?",
     "До обладнання з підвищеною небезпекою належать балони для зріджених, стиснутих і розчинених газів під тиском, вищим за 0,5 бар."),
    ("В яких одиницях виконують розрахунки тиску щодо технічних характеристик запобіжних засобів?",
     "Розрахунки технічних характеристик виконують у міжнародній системі одиниць SI. Для перерахунку використовують співвідношення: 1 бар = 0,1 МПа = 1,019 кгс/см2."),
    ("Яким чином розраховується потужність ударної хвилі вибуху при фізичному вибуху?",
     "Потужність оцінюють через енергію, яка вивільняється при адіабатичному розширенні газу або пари, з урахуванням тривалості вибухового процесу."),
    ("На якому обладнанні застосовують запобіжні клапани?",
     "Запобіжні клапани застосовують на резервуарах, котлах, посудинах, ємностях і трубопроводах для автоматичного скидання надлишкового тиску."),
    ("Які запобіжні клапани найбільш широко використовуються?",
     "Найбільш поширеними є важільно-вантажні та пружинні запобіжні клапани."),
    ("За яких умов клапан вважається малопідйомним?",
     "Клапан вважають малопідйомним, якщо відношення висоти підйому до внутрішнього діаметра тарілки H = h/d не перевищує 0,05."),
    ("Які переваги малопідйомних клапанів перед повнопідйомними?",
     "Малопідйомні клапани мають пропорційну характеристику відкривання, можуть забезпечувати фактичну аварійну витрату і застосовуватися для рідких та газоподібних середовищ."),
    ("Що являє собою пропускна здатність клапана?",
     "Пропускна здатність клапана - це кількість робочого середовища у масових або об'ємних одиницях, яку клапан скидає за одиницю часу при заданих тиску, температурі та ході золотника."),
]


SOURCES = [
    "Методичні вказівки до практичних робіт «Розрахунки запобіжних клапанів» з дисципліни «Безпека праці та професійної діяльності» / уклад. І. О. Мезенцева, С. О. Вамболь, О. О. Кузьменко. Харків: НТУ «ХПІ», 2024. 24 с.",
    "НПАОП 0.00-1.81-18 Правила охорони праці під час експлуатації обладнання, що працює під тиском. Затверджено наказом Міністерства соціальної політики України від 05.03.2018 р. № 333.",
    "ДСТУ ГОСТ 12.2.085:2007 Посудини, що працюють під тиском. Клапани запобіжні. Вимоги щодо безпеки.",
    "Голубенко О. Л., Касьянов М. А., Гунченко О. М. Охорона праці в машинобудівному виробництві. Луганськ: Вид-во Східноукр. ун-ту ім. В. Даля, 2010. 456 с.",
]


def build_docx():
    doc = Document()
    setup_doc(doc)

    for line in [
        "Міністерство освіти і науки України",
        "Національний технічний університет «Харківський політехнічний інститут»",
        "Кафедра безпеки праці та навколишнього середовища",
    ]:
        paragraph(doc, line, align="center", first_line=False)

    for _ in range(5):
        paragraph(doc, "", align="center", first_line=False)

    paragraph(doc, "Практична робота", bold=True, align="center", first_line=False)
    paragraph(doc, "«Розрахунки запобіжних клапанів»", bold=True, align="center", first_line=False)
    paragraph(doc, "з дисципліни «Безпека праці та професійної діяльності»", align="center", first_line=False)
    paragraph(doc, "Варіант 4", bold=True, align="center", first_line=False)

    for _ in range(5):
        paragraph(doc, "", align="center", first_line=False)

    paragraph(doc, "Виконав: студент групи ____________", align="right", first_line=False)
    paragraph(doc, "ПІБ ______________________________", align="right", first_line=False)
    paragraph(doc, "Перевірив: ________________________", align="right", first_line=False)

    for _ in range(5):
        paragraph(doc, "", align="center", first_line=False)
    paragraph(doc, "Харків 2026", align="center", first_line=False)
    doc.add_page_break()

    heading(doc, "Зміст")
    for item in [
        "Вступ",
        "1. Основні теоретичні положення",
        "2. Вихідні дані варіанта 4",
        "3. Розрахунок кількості запобіжних клапанів для котла",
        "4. Підбір запобіжного клапана для пневмопривода",
        "5. Відповіді на контрольні запитання",
        "Висновки",
        "Список використаних джерел",
    ]:
        paragraph(doc, item + " ........................................................................", first_line=False)
    doc.add_page_break()

    heading(doc, "Вступ")
    paragraph(doc, "Мета практичної роботи - ознайомитися з методикою розрахунку запобіжних клапанів як одного із засобів безпечної експлуатації посудин та систем, що працюють під тиском.")
    paragraph(doc, "Необхідність виконання таких розрахунків пов'язана з попередженням аварій, вибухів і виробничого травматизму на виробництвах, де застосовуються котли, резервуари, трубопроводи, пневмогідравлічні системи та інше обладнання під тиском.")

    heading(doc, "1. Основні теоретичні положення")
    paragraph(doc, "Запобіжні клапани призначені для автоматичного випуску рідких, газоподібних середовищ або пари із системи високого тиску при перевищенні допустимого значення. При зниженні тиску до тиску зворотної посадки клапан закривається, і скидання середовища припиняється.")
    paragraph(doc, "Тип клапана визначають за відношенням висоти підйому клапана до внутрішнього діаметра тарілки: H = h/d. Якщо H <= 0,05, клапан є малопідйомним; якщо 0,05 <= H <= 0,25, клапан є повнопідйомним.")
    paragraph(doc, "Пропускна здатність і кількість запобіжних клапанів повинні забезпечувати скидання надлишкового робочого середовища без небезпечного підвищення тиску у захищеній системі.")

    heading(doc, "2. Вихідні дані варіанта 4")
    caption(doc, "Таблиця 1 - Вихідні дані для розрахунку")
    add_table(doc, ["Показник", "Позначення", "Значення", "Одиниця"], [
        ["Паропродуктивність котла", "Gк", G_BOILER, "кг/год"],
        ["Абсолютний тиск пари в котлі", "p", fmt(P_BOILER_MPA, 2), "МПа"],
        ["Висота підйому клапана", "h", fmt(H_LIFT, 2), "см"],
        ["Внутрішній діаметр тарілки клапана", "d", fmt(D_PLATE, 1), "см"],
        ["Надмірний робочий тиск на вході в пневмопривід", "p1р", P1R, "кгс/см2"],
        ["Надлишковий тиск за клапаном", "p2", fmt(P2, 1), "кгс/см2"],
        ["Надлишкова витрата повітря", "G", G_AIR, "кг/год"],
        ["Щільність повітря", "ρ", fmt(RHO, 1), "кг/м3"],
        ["Показник адіабати", "K", fmt(ADIABATIC_K, 1), "-"],
    ])

    heading(doc, "3. Розрахунок кількості запобіжних клапанів для котла")
    paragraph(doc, "Визначаємо тип клапана за формулою:")
    paragraph(doc, "H = h / d.", align="center", first_line=False)
    paragraph(doc, f"H = {fmt(H_LIFT, 2)} / {fmt(D_PLATE, 1)} = {fmt(H_RATIO, 2)}.")
    paragraph(doc, "Оскільки 0,05 <= H <= 0,25, клапан належить до повнопідйомних. Для повнопідйомних клапанів приймаємо коефіцієнт k = 0,015.")
    paragraph(doc, "Кількість запобіжних клапанів визначаємо за формулою:")
    paragraph(doc, "n = Gк / (k · p · d · h).", align="center", first_line=False)
    paragraph(doc, f"p = {fmt(P_BOILER_MPA, 2)} МПа = {int(P_BOILER_PA)} Па.")
    paragraph(doc, f"n = 1300 / (0,015 · 1 640 000 · 4,0 · 0,40) = {fmt(N_CALC, 3)}.")
    paragraph(doc, f"Розрахунково достатньо одного клапана, але при паропродуктивності котла Gк = {G_BOILER} кг/год > 100 кг/год методика вимагає встановлювати не менше двох запобіжних клапанів. Отже, приймаємо n = {N_ACCEPTED} клапани.", bold=True)

    heading(doc, "4. Підбір запобіжного клапана для пневмопривода")
    paragraph(doc, f"Для варіанта 4: p1р = {P1R} кгс/см2, p2 = {fmt(P2, 1)} кгс/см2, G = {G_AIR} кг/год, ρ = {fmt(RHO, 1)} кг/м3, K = {fmt(ADIABATIC_K, 1)}.")
    paragraph(doc, "Робочий тиск належить до діапазону 0,3 МПа < p1р <= 6 МПа, тому максимальний тиск перед клапаном:")
    paragraph(doc, f"p1 = 1,15 · p1р = 1,15 · 26 = {fmt(P1, 2)} кгс/см2.")
    paragraph(doc, f"Відношення тисків: p2/p1 = {fmt(P2, 1)} / {fmt(P1, 2)} = {fmt(P_RATIO, 3)}.")
    paragraph(doc, f"За таблицею для p2/p1 ≈ 0,08 та K = 1,4 приймаємо коефіцієнт B = {fmt(B_COEFF, 3)}. Коефіцієнт витрати клапана приймаємо α = {fmt(ALPHA, 1)}.")
    paragraph(doc, "Необхідна площа перетину клапана визначається з формули G = 0,00159 · α · F · B · √(ρ(p1 - p2)):")
    paragraph(doc, "F = G / [0,00159 · α · B · √(ρ(p1 - p2))].", align="center", first_line=False)
    paragraph(doc, f"F = 21 / [0,00159 · 0,6 · 0,505 · √(14,8 · ({fmt(P1,2)} - {fmt(P2,1)}))] = {fmt(F_NEEDED, 2)} мм2.")
    paragraph(doc, f"За таблицею характеристик запобіжних клапанів обираємо найближчий більший типорозмір: F = {F_SELECTED} мм2, діаметр сідла клапана dc = {SEAT_DIAMETER} мм, діаметр вхідного клапана d1 = {INLET_DIAMETER} мм.", bold=True)

    heading(doc, "5. Відповіді на контрольні запитання")
    for i, (question, answer) in enumerate(CONTROL_ANSWERS, 1):
        paragraph(doc, f"{i}. {question}", bold=True)
        paragraph(doc, answer)

    doc.add_page_break()
    heading(doc, "Висновки")
    paragraph(doc, f"У практичній роботі виконано розрахунок запобіжних клапанів за варіантом 4. Для котла з паропродуктивністю {G_BOILER} кг/год, тиском {fmt(P_BOILER_MPA,2)} МПа, висотою підйому {fmt(H_LIFT,2)} см і діаметром тарілки {fmt(D_PLATE,1)} см визначено H = {fmt(H_RATIO,2)}, тому клапан належить до повнопідйомних.")
    paragraph(doc, f"Розрахункова кількість клапанів становить {fmt(N_CALC,3)}, але з урахуванням вимоги встановлення не менше двох клапанів при паропродуктивності понад 100 кг/год прийнято {N_ACCEPTED} запобіжні клапани.")
    paragraph(doc, f"Для вихідної магістралі пневмопривода визначено максимальний тиск перед клапаном p1 = {fmt(P1,2)} кгс/см2 і необхідну площу перетину F = {fmt(F_NEEDED,2)} мм2. Обрано нормалізований клапан з F = {F_SELECTED} мм2, dc = {SEAT_DIAMETER} мм, d1 = {INLET_DIAMETER} мм.")

    heading(doc, "Список використаних джерел")
    for i, source in enumerate(SOURCES, 1):
        paragraph(doc, f"{i}. {source}", first_line=False)

    doc.save(DOCX_PATH)


def build_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Варіант 4"
    bold = Font(name="Times New Roman", bold=True)
    normal = Font(name="Times New Roman")
    fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = "Розрахунки запобіжних клапанів, варіант 4"
    ws["A1"].font = Font(name="Times New Roman", bold=True, size=14)
    ws.merge_cells("A1:D1")
    ws["A1"].alignment = Alignment(horizontal="center")

    rows = [
        ["Показник", "Формула / значення", "Результат"],
        ["H", "=0.40/4.0", H_RATIO],
        ["Тип клапана", "0,05 <= H <= 0,25", "повнопідйомний"],
        ["n розрах.", "=1300/(0.015*1640000*4.0*0.40)", N_CALC],
        ["n прийняте", "Gк > 100 кг/год", N_ACCEPTED],
        ["p1", "=1.15*26", P1],
        ["p2/p1", "=2.4/p1", P_RATIO],
        ["B", "таблиця 4, K=1,4, p2/p1≈0,08", B_COEFF],
        ["F потрібна, мм2", "=21/(0.00159*0.6*0.505*SQRT(14.8*(29.9-2.4)))", F_NEEDED],
        ["F обрана, мм2", "таблиця 5", F_SELECTED],
        ["dc, мм", "таблиця 5", SEAT_DIAMETER],
        ["d1, мм", "таблиця 5", INLET_DIAMETER],
    ]

    for r, row in enumerate(rows, 3):
        for c, value in enumerate(row, 1):
            cell = ws.cell(r, c, value)
            cell.font = bold if r == 3 else normal
            cell.border = border
            cell.alignment = Alignment(vertical="center")
            if r == 3:
                cell.fill = fill
    for c in range(1, 4):
        ws.column_dimensions[get_column_letter(c)].width = [26, 46, 18][c - 1]

    wb.save(XLSX_PATH)


def build_md():
    lines = [
        "# Практична робота: Розрахунки запобіжних клапанів",
        "",
        "**Варіант:** 4",
        "",
        "## Результати",
        f"- H = {fmt(H_RATIO,2)}; тип клапана: повнопідйомний.",
        f"- Розрахункова кількість клапанів n = {fmt(N_CALC,3)}; прийнято {N_ACCEPTED} клапани.",
        f"- p1 = {fmt(P1,2)} кгс/см2; p2/p1 = {fmt(P_RATIO,3)}; B = {fmt(B_COEFF,3)}.",
        f"- Необхідна площа F = {fmt(F_NEEDED,2)} мм2.",
        f"- Обраний клапан: F = {F_SELECTED} мм2, dc = {SEAT_DIAMETER} мм, d1 = {INLET_DIAMETER} мм.",
    ]
    MD_PATH.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    build_docx()
    build_xlsx()
    build_md()
    print(DOCX_PATH)
    print(XLSX_PATH)
    print(MD_PATH)
